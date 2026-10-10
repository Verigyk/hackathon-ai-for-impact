"""Inspect the two authorized SNCF sources without modifying them."""
from pathlib import Path
from collections import Counter
import csv
import json
import zipfile
import xml.etree.ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / 'data' / 'raw'
OUT = ROOT / 'data' / 'processed'
OUT.mkdir(parents=True, exist_ok=True)

csv_path = RAW / 'Data_Hackathon.csv'
raw = csv_path.read_bytes()
for encoding in ('utf-8-sig', 'cp1252'):
    try:
        content = raw.decode(encoding)
        break
    except UnicodeDecodeError:
        continue
rows = list(csv.DictReader(content.splitlines(), delimiter=';'))
report = {'csv': {
    'encoding': encoding,
    'rows': len(rows),
    'columns': [key for key in rows[0] if key],
    'years': sorted({r['Année'] for r in rows}),
    'months': dict(Counter(r['Mois'] for r in rows)),
    'unique_trains': len({r['N° Train'] for r in rows}),
    'duplicate_train_month_keys': len(rows) - len({(r['Année'], r['Mois'], r['N° Train']) for r in rows}),
    'missing_by_column': {key: sum(not r.get(key, '').strip() for r in rows) for key in rows[0] if key},
}}
xlsx_path = RAW / 'Detail_desserte.xlsx'
if xlsx_path.exists() and zipfile.is_zipfile(xlsx_path):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    report['xlsx'] = []
    for sheet in wb:
        first_rows = list(sheet.iter_rows(min_row=1, max_row=8, values_only=True))
        report['xlsx'].append({'sheet': sheet.title, 'rows': sheet.max_row, 'columns': sheet.max_column, 'first_rows': first_rows})
    wb.close()
else:
    report['xlsx'] = {'status': 'not downloaded or not a valid XLSX'}
(OUT / 'inspection.json').write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
print(json.dumps(report, ensure_ascii=False, default=str))
